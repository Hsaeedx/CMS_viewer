"""
make_earle_metrics.py
Earle-benchmark end-of-life care metrics for the HNC+ICI decedent cohort.

Sheet 1 (Overall): all four metrics stratified by hospice status
Sheet 2 (Risk diff): hospice - no-hospice risk differences with Wald 95% CI

Metrics (all measured in the 30-day window before death_dt):
  - Any ED visit (outpatient rev 0450-0459 OR admitted-from-ED, TYPE_ADM='1')
  - >=2 ED visits (Earle threshold)
  - Any ICU/CCU stay (inpatient rev 020x/021x on any line whose THRU_DT falls in window)
  - Any inpatient admission (ADMSN_DT in window)

Output: tables/earle_metrics.xlsx
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, r'C:\users\hsaee\desktop\cms_viewer\env\Lib\site-packages')

import duckdb
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from math import sqrt
from pathlib import Path

DB_PATH  = r"F:\CMS\cms_data.duckdb"
OUT_PATH = r"C:\Users\hsaee\Desktop\CMS_viewer\projects\HNC_io_hosp\tables\earle_metrics.xlsx"
SQL_PATH = Path(r"C:\Users\hsaee\Desktop\CMS_viewer\projects\HNC_io_hosp\queries\_scratch_earle_metrics.sql")

# ── Style (matches other tables) ─────────────────────────────────────────────
SCARLET, WHITE = 'BA0C2F', 'FFFFFF'
SECTION_PINK, ALT_PINK = 'F5D0D6', 'F9ECEE'

TITLE_FONT    = Font(name='Times New Roman', bold=True,   size=12, color=SCARLET)
HEADER_FONT   = Font(name='Times New Roman', bold=True,   color=WHITE, size=11)
SECTION_FONT  = Font(name='Times New Roman', bold=True,   size=11, color='7A0820')
BODY_FONT     = Font(name='Times New Roman', size=11)
BODY_BOLD     = Font(name='Times New Roman', size=11, bold=True)
NOTE_FONT     = Font(name='Times New Roman', italic=True, size=11, color='555555')

HEADER_FILL  = PatternFill('solid', fgColor=SCARLET)
SECTION_FILL = PatternFill('solid', fgColor=SECTION_PINK)
ALT_FILL     = PatternFill('solid', fgColor=ALT_PINK)

# ── Run SQL ──────────────────────────────────────────────────────────────────
print(f"Running Earle metrics query...")
con = duckdb.connect(DB_PATH, read_only=True)
df = con.execute(SQL_PATH.read_text()).df()
con.close()

# Split rows
rows = {r['stratum']: r for _, r in df.iterrows()}
overall = rows['OVERALL']
hosp    = rows['HOSPICE']
noh     = rows['NO HOSPICE']

# Risk-difference helper
def wald(x1, n1, x2, n2):
    p1, p2 = x1/n1, x2/n2
    rd = (p1 - p2) * 100
    se = sqrt(p1*(1-p1)/n1 + p2*(1-p2)/n2)
    z = 1.959964
    return rd, (p1 - p2 - z*se) * 100, (p1 - p2 + z*se) * 100

METRICS = [
    ('Any ED visit',              'n_any_ed'),
    ('>=2 ED visits',             'n_ge2_ed'),
    ('ICU/CCU stay',              'n_icu'),
    ('Any inpatient admission',   'n_adm'),
]

# ── Build workbook ───────────────────────────────────────────────────────────
print(f"Writing {OUT_PATH} ...")
wb = openpyxl.Workbook()

# ═══════ Sheet 1: Overall + stratified ═══════════════════════════════════════
ws = wb.active
ws.title = 'Overall'

ws.merge_cells('A1:D1')
ws['A1'] = 'Table X. End-of-Life Aggressive Care Metrics (Last 30 Days of Life)'
ws['A1'].font = TITLE_FONT
ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[1].height = 22

# Header row
HEADER_ROW = 3
for ci, h in enumerate(['Metric', f'Overall (n={int(overall["n"]):,})',
                         f'Hospice (n={int(hosp["n"]):,})',
                         f'No hospice (n={int(noh["n"]):,})'], 1):
    cell = ws.cell(row=HEADER_ROW, column=ci, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                               vertical='center', wrap_text=True)
ws.row_dimensions[HEADER_ROW].height = 32

def fmt_np(n, p):
    return f"{int(n):,} ({p:.1f}%)"

r = HEADER_ROW + 1
for idx, (label, col) in enumerate(METRICS):
    pct_col = 'pct_' + col.replace('n_', '')
    vals = [
        label,
        fmt_np(overall[col], overall[pct_col]),
        fmt_np(hosp[col],    hosp[pct_col]),
        fmt_np(noh[col],     noh[pct_col]),
    ]
    fill = ALT_FILL if idx % 2 == 1 else None
    for ci, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=ci, value=v)
        cell.font = BODY_FONT
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                                   vertical='center')
    r += 1

# Footer
r += 1
ws.merge_cells(f'A{r}:D{r}')
note = ws.cell(row=r, column=1, value=(
    'ED = Emergency Department (outpatient revenue codes 0450-0459 or inpatient stays with '
    'admission type = Emergency); ICU/CCU = inpatient revenue codes 0200-0219. '
    '"Any ED" and ">=2 ED" count distinct encounter dates; admitted-from-ED stays are combined '
    'with same-day outpatient ED visits to avoid double-counting. All windows measured in the '
    '30 days ending on the date of death. Denominators shown in column headers.'
))
note.font = NOTE_FONT
note.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[r].height = 65

ws.column_dimensions['A'].width = 34
for c in ['B', 'C', 'D']:
    ws.column_dimensions[c].width = 22

# ═══════ Sheet 2: Risk differences ═══════════════════════════════════════════
ws2 = wb.create_sheet('Risk differences')

ws2.merge_cells('A1:E1')
ws2['A1'] = 'Risk Differences: Hospice vs Non-Enrolled (Last 30 Days of Life)'
ws2['A1'].font = TITLE_FONT
ws2['A1'].alignment = Alignment(horizontal='left', vertical='center')
ws2.row_dimensions[1].height = 22

HEADER_ROW = 3
for ci, h in enumerate([
    'Metric', 'Hospice %', 'No hospice %',
    'Risk difference (pp)', '95% CI (Wald)'
], 1):
    cell = ws2.cell(row=HEADER_ROW, column=ci, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                               vertical='center', wrap_text=True)
ws2.row_dimensions[HEADER_ROW].height = 32

n1, n2 = int(hosp['n']), int(noh['n'])
r = HEADER_ROW + 1
for idx, (label, col) in enumerate(METRICS):
    pct_col = 'pct_' + col.replace('n_', '')
    rd, lo, hi = wald(int(hosp[col]), n1, int(noh[col]), n2)
    ci_str = f"({lo:+.1f} to {hi:+.1f})"
    ci_excl_zero = (lo > 0 and hi > 0) or (lo < 0 and hi < 0)
    vals = [
        label,
        f"{hosp[pct_col]:.1f}",
        f"{noh[pct_col]:.1f}",
        f"{rd:+.1f}",
        ci_str,
    ]
    fill = ALT_FILL if idx % 2 == 1 else None
    for ci, v in enumerate(vals, 1):
        cell = ws2.cell(row=r, column=ci, value=v)
        # Bold the RD + CI columns when CI excludes zero
        if ci in (4, 5) and ci_excl_zero:
            cell.font = BODY_BOLD
        else:
            cell.font = BODY_FONT
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                                   vertical='center')
    r += 1

# Footer w/ caveat
r += 1
ws2.merge_cells(f'A{r}:E{r}')
note = ws2.cell(row=r, column=1, value=(
    'Risk difference = hospice % - no-hospice % (percentage points). Wald 95% confidence '
    'interval on the difference of proportions. Bolded entries: 95% CIs that exclude zero. '
    'Caution: hospice election and terminal-phase inpatient care are near-mutually exclusive '
    'by design (Medicare Hospice Benefit requires forgoing cancer-directed therapy), so the '
    'observed differences are heavily influenced by this administrative structure and the '
    'associated immortal-time bias. Report as descriptive, not causal. A time-to-event / '
    'competing-risks framework (Fine-Gray subdistribution hazard) is the appropriate '
    'confirmatory analysis.'
))
note.font = NOTE_FONT
note.alignment = Alignment(wrap_text=True, vertical='top')
ws2.row_dimensions[r].height = 100

ws2.column_dimensions['A'].width = 30
ws2.column_dimensions['B'].width = 14
ws2.column_dimensions['C'].width = 16
ws2.column_dimensions['D'].width = 22
ws2.column_dimensions['E'].width = 26

wb.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")

from utils import export_xlsx_to_png
FIGURES_DIR = r"C:\Users\hsaee\Desktop\CMS_viewer\projects\HNC_io_hosp\figures"
export_xlsx_to_png(OUT_PATH, FIGURES_DIR)
