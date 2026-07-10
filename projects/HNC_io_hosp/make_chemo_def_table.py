"""
make_chemo_def_table.py
One-off table for PI: how does expanding the chemo-ICI definition
(currently platinum-only) to include docetaxel, paclitaxel, cetuximab, and 5-FU
change the number of patients classified as chemo-ICI vs. ICI monotherapy?

Reports both primary cohort (N=2,527) and broadened cohort (N=5,302).

Output: tables/chemo_def_expansion.xlsx
"""
import sys
sys.path.insert(0, r'C:\users\hsaee\desktop\cms_viewer\env\Lib\site-packages')

import duckdb
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DB_PATH  = r"F:\CMS\cms_data.duckdb"
OUT_PATH = r"C:\Users\hsaee\Desktop\CMS_viewer\projects\HNC_io_hosp\tables\chemo_def_expansion.xlsx"

# Style — matches other tables (Times New Roman; titles 12pt, body 11pt)
SCARLET      = 'BA0C2F'
WHITE        = 'FFFFFF'
SECTION_PINK = 'F5D0D6'
ALT_PINK     = 'F9ECEE'

TITLE_FONT   = Font(name='Times New Roman', bold=True, size=12, color=SCARLET)
HEADER_FONT  = Font(name='Times New Roman', bold=True, color='FFFFFF', size=11)
SECTION_FONT = Font(name='Times New Roman', bold=True, size=11, color='7A0820')
BODY_FONT    = Font(name='Times New Roman', size=11)
NOTE_FONT    = Font(name='Times New Roman', italic=True, size=11, color='555555')

HEADER_FILL  = PatternFill('solid', fgColor=SCARLET)
SECTION_FILL = PatternFill('solid', fgColor=SECTION_PINK)
ALT_FILL     = PatternFill('solid', fgColor=ALT_PINK)


QUERY_TEMPLATE = """
WITH chemo_claims AS (
    SELECT l.DSYSRTKY,
           TRY_STRPTIME(l.THRU_DT,'%Y%m%d') AS chemo_dt,
           l.HCPCS_CD
    FROM io_car_lines l
    JOIN {cohort_table} c ON l.DSYSRTKY = c.DSYSRTKY
    WHERE l.HCPCS_CD IN ('J9060','J9045','J9171','J9267','J9055','J9190')

    UNION ALL

    SELECT r.DSYSRTKY,
           TRY_STRPTIME(COALESCE(NULLIF(r.REV_DT,''), r.THRU_DT),'%Y%m%d') AS chemo_dt,
           r.HCPCS_CD
    FROM io_out_revenue r
    JOIN {cohort_table} c ON r.DSYSRTKY = c.DSYSRTKY
    WHERE r.HCPCS_CD IN ('J9060','J9045','J9171','J9267','J9055','J9190')
),
chemo_ici_matches AS (
    SELECT DISTINCT
        ch.DSYSRTKY,
        CASE
            WHEN ch.HCPCS_CD IN ('J9060','J9045') THEN 'platinum'
            WHEN ch.HCPCS_CD = 'J9055'            THEN 'cetuximab'
            WHEN ch.HCPCS_CD = 'J9171'            THEN 'docetaxel'
            WHEN ch.HCPCS_CD = 'J9267'            THEN 'paclitaxel'
            WHEN ch.HCPCS_CD = 'J9190'            THEN '5-FU'
        END AS agent_grp
    FROM chemo_claims ch
    JOIN io_claims_raw ici ON ch.DSYSRTKY = ici.DSYSRTKY
    WHERE ABS(datediff('day', ch.chemo_dt, ici.io_date)) <= 21
),
pt_flags AS (
    SELECT
        c.DSYSRTKY,
        BOOL_OR(m.agent_grp = 'platinum')   AS w_plat,
        BOOL_OR(m.agent_grp = 'cetuximab')  AS w_cetux,
        BOOL_OR(m.agent_grp = 'docetaxel')  AS w_doce,
        BOOL_OR(m.agent_grp = 'paclitaxel') AS w_pacli,
        BOOL_OR(m.agent_grp = '5-FU')       AS w_5fu
    FROM {cohort_table} c
    LEFT JOIN chemo_ici_matches m ON c.DSYSRTKY = m.DSYSRTKY
    GROUP BY c.DSYSRTKY
)
SELECT
    COUNT(*) AS n_cohort,
    COUNT(*) FILTER (WHERE w_plat) AS n_current,
    COUNT(*) FILTER (WHERE w_plat OR w_cetux OR w_doce OR w_pacli OR w_5fu) AS n_expanded,
    COUNT(*) FILTER (WHERE w_cetux  AND NOT w_plat) AS add_cetux,
    COUNT(*) FILTER (WHERE w_pacli  AND NOT w_plat) AS add_pacli,
    COUNT(*) FILTER (WHERE w_doce   AND NOT w_plat) AS add_doce,
    COUNT(*) FILTER (WHERE w_5fu    AND NOT w_plat) AS add_5fu,
    COUNT(*) FILTER (WHERE w_cetux)  AS ever_cetux,
    COUNT(*) FILTER (WHERE w_pacli)  AS ever_pacli,
    COUNT(*) FILTER (WHERE w_doce)   AS ever_doce,
    COUNT(*) FILTER (WHERE w_5fu)    AS ever_5fu
FROM pt_flags
"""


def fmt_n_pct(n, total):
    return f"{n:,} ({100.0 * n / total:.1f}%)"


def query(con, cohort_table):
    row = con.execute(QUERY_TEMPLATE.format(cohort_table=cohort_table)).fetchone()
    keys = ['n_cohort','n_current','n_expanded',
            'add_cetux','add_pacli','add_doce','add_5fu',
            'ever_cetux','ever_pacli','ever_doce','ever_5fu']
    return dict(zip(keys, row))


print("Querying both cohorts...")
con = duckdb.connect(DB_PATH, read_only=True)
con.execute("SET memory_limit='24GB'; SET threads=12;")
primary  = query(con, 'io_cohort')
broaden  = query(con, 'io_cohort_itc')
con.close()

for label, d in [('Primary', primary), ('Broadened', broaden)]:
    print(f"  {label}: N={d['n_cohort']:,}  current={d['n_current']:,}  expanded={d['n_expanded']:,}  "
          f"add={d['n_expanded']-d['n_current']:,}")


# ── Build workbook ────────────────────────────────────────────────────────────
print(f"Writing {OUT_PATH} ...")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Chemo-ICI Expansion'

# Title
ws.merge_cells('A1:E1')
ws['A1'] = ('Chemo-Immunotherapy Definition: Impact of Adding Non-Platinum Agents '
            '(docetaxel, paclitaxel, cetuximab, 5-FU)')
ws['A1'].font = TITLE_FONT
ws['A1'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 32

# Subtitle
ws.merge_cells('A2:E2')
ws['A2'] = ('Patients classified as chemo-ICI under current definition (platinum only) vs. '
            'expanded definition. Chemo-ICI = any listed agent within ±21 days of any ICI dose.')
ws['A2'].font = Font(name='Times New Roman', italic=True, size=11, color='555555')
ws['A2'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws.row_dimensions[2].height = 30

# Header
HEADERS = ['Metric', 'Primary Cohort', '% of cohort', 'Broadened Cohort', '% of cohort']
HEADER_ROW = 4
for ci, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=HEADER_ROW, column=ci, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                               vertical='center', wrap_text=True)
ws.row_dimensions[HEADER_ROW].height = 26


def write_row(r, label, primary_n, broaden_n, *,
              section=False, indent=False, alt=False, bold_delta=False):
    n_p = primary['n_cohort']
    n_b = broaden['n_cohort']

    if section:
        ws.merge_cells(f'A{r}:E{r}')
        cell = ws.cell(row=r, column=1, value=label)
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL
        cell.alignment = Alignment(horizontal='left', vertical='center')
        return

    vals = [
        label,
        f"{primary_n:,}",
        f"{100.0*primary_n/n_p:.1f}%",
        f"{broaden_n:,}",
        f"{100.0*broaden_n/n_b:.1f}%",
    ]
    for ci, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=ci, value=v)
        cell.font = Font(name='Times New Roman', size=11, bold=bold_delta)
        if alt:
            cell.fill = ALT_FILL
        cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                                   vertical='center',
                                   indent=2 if (indent and ci == 1) else 0)


r = HEADER_ROW + 1
write_row(r, 'COHORT SIZE', None, None, section=True)
r += 1
write_row(r, 'Total patients', primary['n_cohort'], broaden['n_cohort'])
r += 2

write_row(r, 'CHEMO-ICI CLASSIFICATION', None, None, section=True)
r += 1
write_row(r, 'Current definition (platinum only: J9060, J9045)',
          primary['n_current'], broaden['n_current'], alt=False)
r += 1
write_row(r, 'Expanded definition (+ J9171, J9267, J9055, J9190)',
          primary['n_expanded'], broaden['n_expanded'], alt=True)
r += 1
write_row(r, 'Net newly captured by expansion',
          primary['n_expanded'] - primary['n_current'],
          broaden['n_expanded']  - broaden['n_current'],
          bold_delta=True)
r += 2

write_row(r, 'PATIENTS NEWLY CAPTURED BY EACH AGENT (within ±21 days of ICI, in patients with no platinum within ±21 days of ICI)',
          None, None, section=True)
r += 1
for idx, (label, key) in enumerate([
    ('Cetuximab (J9055)',  'add_cetux'),
    ('Paclitaxel (J9267)', 'add_pacli'),
    ('Docetaxel (J9171)',  'add_doce'),
    ('5-Fluorouracil (J9190)', 'add_5fu'),
]):
    write_row(r, label, primary[key], broaden[key],
              indent=True, alt=(idx % 2 == 1))
    r += 1
r += 1

write_row(r, 'PATIENTS EVER RECEIVING EACH AGENT WITHIN ±21 DAYS OF ICI (includes patients also on platinum — overlap with current definition)',
          None, None, section=True)
r += 1
for idx, (label, key) in enumerate([
    ('Cetuximab (J9055)',  'ever_cetux'),
    ('Paclitaxel (J9267)', 'ever_pacli'),
    ('Docetaxel (J9171)',  'ever_doce'),
    ('5-Fluorouracil (J9190)', 'ever_5fu'),
]):
    write_row(r, label, primary[key], broaden[key],
              indent=True, alt=(idx % 2 == 1))
    r += 1
r += 2

# Footnote
ws.merge_cells(f'A{r}:E{r}')
note = ws.cell(row=r, column=1, value=(
    'Primary cohort: HNC + ICI patients with documented prior curative-intent therapy and ≥180 days from diagnosis to ICI start. '
    'Broadened cohort: all eligible HNC + ICI patients regardless of prior curative therapy or dx-to-ICI interval. '
    'Chemo-ICI flag requires a chemo claim (J-code) within ±21 days of any pembrolizumab (J9271) or nivolumab (J9299) administration. '
    '"Newly captured" = patients reclassified from ICI monotherapy → chemo-ICI under the expanded definition; '
    'these patients had at least one of the four added agents within ±21 days of ICI and did NOT have platinum within ±21 days of ICI. '
    'Patients with both platinum and a new agent are already captured under the current definition.'
))
note.font = NOTE_FONT
note.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[r].height = 70

# Column widths
ws.column_dimensions['A'].width = 70
for col in ['B', 'C', 'D', 'E']:
    ws.column_dimensions[col].width = 18

wb.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
