"""
make_trend_test.py
Formal secular-trend test for each binary outcome across the study window.

For each outcome, fits logistic regression with death_year as a continuous
predictor:
  logit(P(outcome)) = beta_0 + beta_1 * (death_year - 2017)

Reports OR per year + Wald 95% CI + p-value.
Unadjusted and adjusted-for-subsite+regimen versions.

Format per PI: "No significant temporal trend in [outcome] across 2017-2023
(OR per year X.XX, 95% CI X-X)."

Output: tables/trend_test.xlsx
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, r'C:\users\hsaee\desktop\cms_viewer\env\Lib\site-packages')

import duckdb
import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

DB_PATH  = r"F:\CMS\cms_data.duckdb"
OUT_PATH = r"C:\Users\hsaee\Desktop\CMS_viewer\projects\HNC_io_hosp\tables\trend_test.xlsx"

# ── Load data ────────────────────────────────────────────────────────────────
print("Loading io_analytic ...")
con = duckdb.connect(DB_PATH, read_only=True)
df = con.execute("""
    SELECT
        hospice_enrolled, hospice_short_stay,
        in_hospital_death,
        io_within_14d_of_death, io_within_30d_of_death,
        any_ed_last_30d, ge_2_ed_last_30d,
        icu_last_30d, admission_last_30d,
        death_year, subsite_category, io_regimen
    FROM io_analytic
""").df()
con.close()

df['death_year_c'] = df['death_year'].astype(float) - 2017
df['io_regimen']   = df['io_regimen'].replace({'IO monotherapy': 'ICI monotherapy',
                                                'chemo-IO':       'chemo-ICI'})

# Fill outcome NaNs to 0 where appropriate
for c in ['hospice_enrolled', 'in_hospital_death', 'any_ed_last_30d',
          'ge_2_ed_last_30d', 'icu_last_30d', 'admission_last_30d',
          'io_within_14d_of_death', 'io_within_30d_of_death']:
    df[c] = df[c].fillna(0).astype(int)

# ── Outcomes to test ─────────────────────────────────────────────────────────
OUTCOMES = [
    ('Hospice enrolled',                'hospice_enrolled',       None),
    ('In-hospital death',               'in_hospital_death',      None),
    ('ICI within 14 days of death',     'io_within_14d_of_death', None),
    ('ICI within 30 days of death',     'io_within_30d_of_death', None),
    ('Any inpatient admission last 30d','admission_last_30d',     None),
    ('>=2 ED visits last 30d',          'ge_2_ed_last_30d',       None),
    ('Any ED visit last 30d',           'any_ed_last_30d',        None),
    ('ICU/CCU stay last 30d',           'icu_last_30d',           None),
    # hospice_short_stay only defined among enrollees
    ('Hospice short stay (<=7d, among enrollees)', 'hospice_short_stay',
        lambda d: d[d['hospice_enrolled'] == 1].copy()),
]

def fit(data, outcome, formula):
    """Fit logit; return OR/CI/p on the death_year_c term."""
    model = smf.logit(f"{outcome} ~ {formula}", data=data).fit(
        method='newton', maxiter=200, disp=False)
    coef = float(model.params['death_year_c'])
    se   = float(model.bse['death_year_c'])
    p    = float(model.pvalues['death_year_c'])
    or_ = np.exp(coef)
    lo, hi = np.exp(coef - 1.96*se), np.exp(coef + 1.96*se)
    n = int(model.nobs)
    return or_, lo, hi, p, n

rows = []
for label, outcome, subset_fn in OUTCOMES:
    d = subset_fn(df) if subset_fn else df
    # Guard: outcome must have some variation in the subset
    if d[outcome].sum() == 0 or d[outcome].sum() == len(d):
        print(f"  {label}: skipped (no outcome variation)")
        continue

    # Unadjusted
    try:
        or_u, lo_u, hi_u, p_u, n_u = fit(d, outcome, "death_year_c")
    except Exception as e:
        print(f"  {label} unadj: FAILED — {e}")
        or_u = lo_u = hi_u = p_u = n_u = None

    # Adjusted for subsite + regimen
    try:
        or_a, lo_a, hi_a, p_a, n_a = fit(d, outcome,
            "death_year_c + C(subsite_category) + C(io_regimen)")
    except Exception as e:
        print(f"  {label} adj: FAILED — {e}")
        or_a = lo_a = hi_a = p_a = n_a = None

    n_events = int(d[outcome].sum())
    n_total  = int(len(d))
    rows.append({
        'outcome':  label,
        'n_total':  n_total,
        'n_events': n_events,
        'or_u':     or_u, 'lo_u': lo_u, 'hi_u': hi_u, 'p_u': p_u,
        'or_a':     or_a, 'lo_a': lo_a, 'hi_a': hi_a, 'p_a': p_a,
    })

# ── Print summary ────────────────────────────────────────────────────────────
def fmt_or(o, lo, hi):
    if o is None: return ''
    return f"{o:.2f} ({lo:.2f}-{hi:.2f})"

def fmt_p(p):
    if p is None: return ''
    if p < 0.001: return '<0.001'
    return f"{p:.3f}"

print("\n" + "="*100)
print(f"{'Outcome':<44} {'Events/N':<14} {'Unadj OR/yr (95% CI)':<26} {'p':<8} {'Adj OR/yr (95% CI)':<26} {'p':<8}")
print("-"*130)
for r in rows:
    print(f"{r['outcome']:<44} {r['n_events']}/{r['n_total']:<10} "
          f"{fmt_or(r['or_u'], r['lo_u'], r['hi_u']):<26} {fmt_p(r['p_u']):<8} "
          f"{fmt_or(r['or_a'], r['lo_a'], r['hi_a']):<26} {fmt_p(r['p_a']):<8}")

# ── Write Excel ──────────────────────────────────────────────────────────────
print(f"\nWriting {OUT_PATH} ...")

SCARLET, WHITE = 'BA0C2F', 'FFFFFF'
ALT_PINK = 'F9ECEE'

TITLE_FONT   = Font(name='Times New Roman', bold=True, size=12, color=SCARLET)
HEADER_FONT  = Font(name='Times New Roman', bold=True, color=WHITE, size=11)
BODY_FONT    = Font(name='Times New Roman', size=11)
BODY_BOLD    = Font(name='Times New Roman', size=11, bold=True)
NOTE_FONT    = Font(name='Times New Roman', italic=True, size=11, color='555555')
HEADER_FILL  = PatternFill('solid', fgColor=SCARLET)
ALT_FILL     = PatternFill('solid', fgColor=ALT_PINK)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Trend tests'

ws.merge_cells('A1:H1')
ws['A1'] = 'Table X. Secular Trend Tests for End-of-Life Care Outcomes (2017-2023)'
ws['A1'].font = TITLE_FONT
ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[1].height = 22

HEADER_ROW = 3
headers = ['Outcome', 'Events / N',
           'Unadjusted OR per year', 'Unadj 95% CI', 'Unadj p',
           'Adjusted OR per year',   'Adj 95% CI',   'Adj p']
for ci, h in enumerate(headers, 1):
    cell = ws.cell(row=HEADER_ROW, column=ci, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                               vertical='center', wrap_text=True)
ws.row_dimensions[HEADER_ROW].height = 36

r = HEADER_ROW + 1
for idx, row in enumerate(rows):
    def _fmt(v, fmt='.2f'):
        return '' if v is None else format(v, fmt)

    ci_u = f"({_fmt(row['lo_u'])}-{_fmt(row['hi_u'])})" if row['or_u'] else ''
    ci_a = f"({_fmt(row['lo_a'])}-{_fmt(row['hi_a'])})" if row['or_a'] else ''
    vals = [
        row['outcome'],
        f"{row['n_events']:,} / {row['n_total']:,}",
        _fmt(row['or_u']),
        ci_u,
        fmt_p(row['p_u']),
        _fmt(row['or_a']),
        ci_a,
        fmt_p(row['p_a']),
    ]
    fill = ALT_FILL if idx % 2 == 1 else None
    # Bold OR/CI when CI excludes 1 (significant at 0.05 by Wald)
    sig_u = row['or_u'] is not None and (row['lo_u'] > 1 or row['hi_u'] < 1)
    sig_a = row['or_a'] is not None and (row['lo_a'] > 1 or row['hi_a'] < 1)
    for ci, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=ci, value=v)
        bold = (sig_u and ci in (3, 4)) or (sig_a and ci in (6, 7))
        cell.font = BODY_BOLD if bold else BODY_FONT
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                                   vertical='center')
    r += 1

# Footer
r += 1
ws.merge_cells(f'A{r}:H{r}')
note = ws.cell(row=r, column=1, value=(
    'Logistic regression with year of death as a continuous predictor (centered at 2017). '
    'OR per year = multiplicative change in the odds of the outcome for each additional calendar year. '
    'Adjusted models include HNC subsite category and ICI regimen (monotherapy vs chemo-ICI) as fixed effects. '
    'Bolded OR/CI cells indicate 95% CIs that exclude 1.0 (nominally significant temporal trend). '
    'Non-bolded rows: temporal trend not statistically significant across 2017-2023.'
))
note.font = NOTE_FONT
note.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[r].height = 70

ws.column_dimensions['A'].width = 44
ws.column_dimensions['B'].width = 15
for c in ['C','D','E','F','G','H']:
    ws.column_dimensions[c].width = 17

wb.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")

from utils import export_xlsx_to_png
FIGURES_DIR = r"C:\Users\hsaee\Desktop\CMS_viewer\projects\HNC_io_hosp\figures"
export_xlsx_to_png(OUT_PATH, FIGURES_DIR)
