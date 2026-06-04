"""
comp_b_sensitivity.py

Supplementary dose-response analysis: week-by-week SLP timing, each vs Week 5+ reference.
  Week 1 (days  8-14) vs Week 5+ (days 36-90)  — ref
  Week 2 (days 15-21) vs Week 5+ (days 36-90)  — ref
  Week 3 (days 22-28) vs Week 5+ (days 36-90)  — ref
  Week 4 (days 29-35) vs Week 5+ (days 36-90)  — ref

Week 0 (days 1-7) excluded entirely, consistent with primary analysis.
Reference group: Week 5+ (days 36-90), same as primary Early vs Late comparison.
PSM + TV Cox per week bin. Dose-response: earlier SLP → greater benefit.
"""
import numpy as np
import pandas as pd
from scipy.spatial import cKDTree
from sklearn.linear_model import LogisticRegression
from sklearn.preprocessing import StandardScaler
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from stroke_slp_common import OUT_DIR, add_aspiration_related_pna, connect, run_tv_cox

FIGURE_PATH = OUT_DIR / "Supp_Figure2.png"
TABLE_PATH  = OUT_DIR / "Supp_Table2.xlsx"

RANDOM_SEED   = 42
CONT_VARS     = ['age_at_adm', 'index_los', 'van_walraven_score', 'adm_year']
BINARY_VARS   = ['afib', 'hypertension', 'mech_vent', 'prior_stroke', 'dual_eligible']
CAT_VARS      = ['sex', 'race', 'stroke_type', 'drg_group', 'adm_source', 'rucc_group']

# Week-based bins — Week 0 (days 1-7) excluded (not in any bin)
BINS = {
    'Wk1':  ( 8,  14),
    'Wk2':  (15,  21),
    'Wk3':  (22,  28),
    'Wk4':  (29,  35),
    'Wk5+': (36,  90),   # reference
}
COMPARISONS = [
    ('Wk1', 'Wk1',  'Wk5+'),
    ('Wk2', 'Wk2',  'Wk5+'),
    ('Wk3', 'Wk3',  'Wk5+'),
    ('Wk4', 'Wk4',  'Wk5+'),
]
OUTCOMES = [
    ('Aspiration-related PNA', 'days_to_asp_related', 'days_to_death', False),
    ('PEG/G-tube',             'days_to_gtube',       'days_to_death', True),
    ('Mortality',              'days_to_death',        None,            False),
]

def bucket_drg(d):
    if pd.isna(d): return 'Other'
    try: n = int(str(d).strip())
    except: return 'Other'
    if 61 <= n <= 69: return 'Medical_stroke'
    if 20 <= n <= 38: return 'Neurosurgical'
    if 52 <= n <= 60: return 'Spinal'
    if 70 <= n <= 74: return 'TIA_headache'
    return 'Other'


# ── Load ───────────────────────────────────────────────────────────────────────
print("Loading data...")
con = connect(read_only=True)
df_prop = con.execute("""
    SELECT DSYSRTKY, days_to_slp_outpt,
           age_at_adm, index_los, van_walraven_score, adm_year,
           sex, race, stroke_type, DRG_CD AS drg_cd, adm_source,
           mech_vent, peg_placed, prior_stroke, afib, hypertension,
           rucc_group, dual_eligible
    FROM stroke_propensity WHERE days_to_slp_outpt IS NOT NULL
""").df()
df_out = con.execute("""
    SELECT DSYSRTKY, days_to_death, days_to_aspiration, days_to_gtube,
           days_to_pneumonia, first_pneumonia_code, pre_stroke_tube
    FROM stroke_outcomes
""").df()
con.close()

def assign_group(d):
    for grp, (lo, hi) in BINS.items():
        if lo <= d <= hi: return grp
    return None   # days 1-7 (Wk0) and day 0 → None → filtered out

df_prop['timing_new'] = df_prop['days_to_slp_outpt'].apply(assign_group)
df_prop = df_prop[df_prop['timing_new'].notna()].copy()
for col in CONT_VARS:
    df_prop[col] = df_prop[col].fillna(df_prop[col].median())
df_prop['drg_group']    = df_prop['drg_cd'].apply(bucket_drg)
df_prop['adm_source']   = df_prop['adm_source'].fillna('Unknown').astype(str)
df_prop['adm_year']     = df_prop['adm_year'].fillna(df_prop['adm_year'].median()).astype(int)
df_prop['rucc_group']   = df_prop['rucc_group'].fillna('Unknown').astype(str)
df_prop['dual_eligible'] = df_prop['dual_eligible'].fillna(0).astype(int)

df_all = df_prop.merge(df_out, on='DSYSRTKY', how='inner')
df_all = add_aspiration_related_pna(df_all)
counts = df_all['timing_new'].value_counts()
print(f"  {len(df_all):,} rows  |  " +
      "  ".join(f"{g}:{counts.get(g,0):,}" for g in ['Wk1','Wk2','Wk3','Wk4','Wk5+']))


# ── PSM ────────────────────────────────────────────────────────────────────────
def run_psm(df_full, treat_grp, ctrl_grp):
    df = df_full[df_full['timing_new'].isin([treat_grp, ctrl_grp])].copy()
    df['_treated'] = (df['timing_new'] == treat_grp).astype(int)
    n_treat  = df['_treated'].sum()
    dummies  = pd.get_dummies(df[CAT_VARS], drop_first=True)
    X        = pd.concat([df[CONT_VARS + BINARY_VARS].astype(float), dummies], axis=1)
    X_scaled = StandardScaler().fit_transform(X)
    lr = LogisticRegression(max_iter=1000, solver='lbfgs', C=1.0)
    lr.fit(X_scaled, df['_treated'])
    ps       = np.clip(lr.predict_proba(X_scaled)[:, 1], 1e-6, 1 - 1e-6)
    logit_ps = np.log(ps / (1 - ps))
    df = df.reset_index(drop=True)
    df['logit_ps'] = logit_ps
    caliper = 0.2 * logit_ps.std()
    treated = df[df['_treated'] == 1].reset_index(drop=True)
    control = df[df['_treated'] == 0].reset_index(drop=True)
    tree    = cKDTree(control['logit_ps'].values.reshape(-1, 1))
    rng     = np.random.default_rng(RANDOM_SEED)
    matched_pairs, used_ctrl = [], set()
    for i in rng.permutation(len(treated)):
        t_lps = treated.loc[i, 'logit_ps']
        t_id  = treated.loc[i, 'DSYSRTKY']
        dists, idxs = tree.query([[t_lps]], k=min(50, len(control)))
        for dist, idx in zip(dists[0], idxs[0]):
            if dist > caliper: break
            cid = control.loc[idx, 'DSYSRTKY']
            if cid not in used_ctrl:
                matched_pairs.append((t_id, cid))
                used_ctrl.add(cid)
                break
    n_matched = len(matched_pairs)
    print(f"  {treat_grp} vs {ctrl_grp}: {n_matched:,}/{n_treat:,} matched "
          f"({100*n_matched/n_treat:.1f}%) caliper={caliper:.4f}")
    matched_ids = {pid for pair in matched_pairs for pid in pair}
    return matched_ids


# ── TV Cox ─────────────────────────────────────────────────────────────────────
# ── Run all comparisons ────────────────────────────────────────────────────────
results = []
for comp_label, treat_grp, ctrl_grp in COMPARISONS:
    print(f"\n{comp_label}: {treat_grp} vs {ctrl_grp} (ref)")
    matched_ids = run_psm(df_all, treat_grp, ctrl_grp)
    df_comp = df_all[df_all['DSYSRTKY'].isin(matched_ids)].copy()
    for out_label, ev_col, comp_col, excl_peg in OUTCOMES:
        sub = df_comp.copy()
        if excl_peg:
            sub = sub[(sub['peg_placed'] == 0) & (sub['pre_stroke_tube'].fillna(0) == 0)]
        res, n, n_ev = run_tv_cox(sub, ev_col, comp_col, treat_grp, group_col="timing_new")
        entry = {'Comparison': f"{treat_grp} vs {ctrl_grp} (ref)",
                 'comp_label': comp_label, 'treat_grp': treat_grp,
                 'Outcome': out_label, 'N': n, 'Events': n_ev,
                 'Event_pct': round(100 * n_ev / n, 1) if n > 0 else 0}
        if res:
            hr, lo, hi, p = res
            p_str = '<0.0001' if p < 0.0001 else f'{p:.4f}'
            entry.update({'HR': round(hr, 2), 'CI_low': round(lo, 2), 'CI_high': round(hi, 2),
                          'p_raw': p, 'p_str': p_str})
            tag = f"HR={hr:.2f} [{lo:.2f}-{hi:.2f}] p={p_str}"
        else:
            entry.update({'HR': np.nan, 'CI_low': np.nan, 'CI_high': np.nan,
                          'p_raw': np.nan, 'p_str': 'N/A'})
            tag = "failed"
        results.append(entry)
        print(f"  {out_label}: n={n:,} ev={n_ev:,}  {tag}")

df_res = pd.DataFrame(results)

print("\n" + "="*90)
print("WEEK-BY-WEEK DOSE-RESPONSE (Wks 1-4 each vs Wk5+ ref) — PSM + TV Cox")
print("="*90)
print(f"{'Comp':<6} {'Outcome':<18} {'HR [95% CI] p'}")
print("-"*65)
for _, row in df_res.iterrows():
    res_str = (f"HR={row['HR']:.2f} [{row['CI_low']:.2f}-{row['CI_high']:.2f}] p={row['p_str']}"
               if pd.notna(row['HR']) else "failed")
    print(f"{row['comp_label']:<6} {row['Outcome']:<18} {res_str}")


# ── Forest plot — week-by-week dose-response ───────────────────────────────────
PRIMARY   = ['Aspiration-related PNA', 'PEG/G-tube', 'Mortality']

# OSU brand color gradient: scarlet → dark 40 → dark 60
COMP_COL = {
    'Wk1': '#ba0c2f',
    'Wk2': '#94091f',
    'Wk3': '#70071c',
    'Wk4': '#4a0513',
}
COMP_MKR = {'Wk1': 'o', 'Wk2': 's', 'Wk3': 'D', 'Wk4': '^'}
COMP_LBL = {
    'Wk1': 'Week 1 (days 8\u201314)  vs Week 5+ ref',
    'Wk2': 'Week 2 (days 15\u201321) vs Week 5+ ref',
    'Wk3': 'Week 3 (days 22\u201328) vs Week 5+ ref',
    'Wk4': 'Week 4 (days 29\u201335) vs Week 5+ ref',
}
COMP_ROW_LBL = {
    'Wk1': 'Week 1 (days 8\u201314)',
    'Wk2': 'Week 2 (days 15\u201321)',
    'Wk3': 'Week 3 (days 22\u201328)',
    'Wk4': 'Week 4 (days 29\u201335)',
}

# y-positions: 3 outcome sections × 4 week rows + 3 headers
ROW_Y = {
    'asp_hdr':               15.5,
    ('Aspiration-related PNA','Wk1'): 14.65,
    ('Aspiration-related PNA','Wk2'): 13.80,
    ('Aspiration-related PNA','Wk3'): 12.95,
    ('Aspiration-related PNA','Wk4'): 12.10,
    'gtube_hdr':              10.85,
    ('PEG/G-tube','Wk1'):    10.00,
    ('PEG/G-tube','Wk2'):     9.15,
    ('PEG/G-tube','Wk3'):     8.30,
    ('PEG/G-tube','Wk4'):     7.45,
    'mort_hdr':                6.20,
    ('Mortality','Wk1'):       5.35,
    ('Mortality','Wk2'):       4.50,
    ('Mortality','Wk3'):       3.65,
    ('Mortality','Wk4'):       2.80,
}
YLIM = (2.15, 16.75)

fig = plt.figure(figsize=(14, 11), facecolor='white')
gs  = gridspec.GridSpec(1, 3, width_ratios=[3.2, 5.0, 3.5],
                        wspace=0.0, left=0.01, right=0.99, top=0.92, bottom=0.07)
ax_L, ax_M, ax_R = [fig.add_subplot(gs[i]) for i in range(3)]

for ax in [ax_L, ax_M, ax_R]:
    ax.set_ylim(*YLIM)
    ax.set_yticks([])
    ax.set_facecolor('white')
    # Alternating section shading
    ax.axhspan(11.70, 16.25, color='#fdf5f6', alpha=0.85, zorder=0)
    ax.axhspan(7.10,  11.30, color='#f9eaeb', alpha=0.85, zorder=0)
    ax.axhspan(2.45,   6.65, color='#fdf5f6', alpha=0.85, zorder=0)

for ax in [ax_L, ax_R]:
    ax.set_xticks([])
    for sp in ax.spines.values(): sp.set_visible(False)

HDR_KW = dict(fontsize=9, fontweight='bold', color='#555555', va='bottom')
ax_L.text(0.97, YLIM[1] + 0.02, 'Outcome / Week', ha='right',
          transform=ax_L.get_yaxis_transform(), **HDR_KW)
ax_M.axvline(1.0, color='#888888', lw=1.2, linestyle='--', zorder=1)
ax_R.text(0.03, YLIM[1] + 0.02, 'HR [95% CI]              p-value', ha='left',
          transform=ax_R.get_yaxis_transform(), **HDR_KW)


def sec_header(y, text):
    for ax in [ax_L, ax_M, ax_R]:
        ax.axhline(y - 0.28, color='#CCCCCC', lw=0.8, zorder=0)
    ax_L.text(0.97, y, text, ha='right', va='center', fontsize=10.5,
              fontweight='bold', color='#70071c', transform=ax_L.get_yaxis_transform())


sec_header(ROW_Y['asp_hdr'],   'Aspiration-related Pneumonia')
sec_header(ROW_Y['gtube_hdr'], 'PEG / G-tube Placement')
sec_header(ROW_Y['mort_hdr'],  'All-cause Mortality')

for _, row in df_res[df_res['Outcome'].isin(PRIMARY)].iterrows():
    key = (row['Outcome'], row['comp_label'])
    if key not in ROW_Y or pd.isna(row['HR']): continue
    ypos = ROW_Y[key]
    hr, lo, hi = row['HR'], row['CI_low'], row['CI_high']
    col = COMP_COL[row['comp_label']]
    mkr = COMP_MKR[row['comp_label']]
    lbl = COMP_ROW_LBL[row['comp_label']]
    ax_L.text(0.97, ypos, f"  {lbl}", ha='right', va='center',
              fontsize=9, color='#333333', transform=ax_L.get_yaxis_transform())
    ax_M.plot([lo, hi], [ypos, ypos], color=col, lw=2.0, solid_capstyle='round', zorder=3)
    ax_M.plot(hr, ypos, marker=mkr, color=col, markersize=9,
              markeredgecolor='white', markeredgewidth=0.8, zorder=4)
    p_fmt = row['p_str'] if row['p_str'].startswith('<') else f"= {row['p_str']}"
    ax_R.text(0.03, ypos, f"  {hr:.2f}  [{lo:.2f}\u2013{hi:.2f}]     p {p_fmt}",
              ha='left', va='center', fontsize=9, color='#70071c',
              transform=ax_R.get_yaxis_transform())

ax_M.set_xscale('log')
ax_M.set_xlim(0.30, 2.8)
ax_M.set_xticks([0.33, 0.5, 0.7, 1.0, 1.25, 1.5, 2.0, 2.5])
ax_M.set_xticklabels(['0.33', '0.50', '0.70', '1.00', '1.25', '1.50', '2.00', '2.50'],
                     fontsize=8.5)
ax_M.set_xlabel('Hazard Ratio (log scale)\nReference: Late SLP (Week 5+, days 36\u201390)',
                fontsize=9.5, color='#333333')
for sp in ['top', 'left', 'right']:
    ax_M.spines[sp].set_visible(False)
ax_M.spines['bottom'].set_color('#AAAAAA')
ax_M.tick_params(axis='x', length=3, color='#AAAAAA')

from matplotlib.lines import Line2D
ax_M.legend(
    handles=[
        Line2D([0],[0], marker=COMP_MKR[k], color='w', markerfacecolor=COMP_COL[k],
               markersize=9, markeredgecolor='white', label=COMP_LBL[k])
        for k in ['Wk1', 'Wk2', 'Wk3', 'Wk4']
    ],
    loc='lower right', fontsize=8.5, frameon=True, framealpha=0.95,
    edgecolor='#CCCCCC', title='SLP Week (vs Week 5+ ref)',
    title_fontsize=8.5,
)

fig.suptitle(
    'Supplementary: Week-by-Week Dose-Response \u2014 SLP Timing and Post-Stroke Outcomes\n'
    'PSM-matched (each week vs Week 5+ ref)  \u2022  Time-Varying Cox  \u2022  '
    'Max follow-up 365 days  \u2022  Week 0 (days 1\u20137) excluded',
    fontsize=11, fontweight='bold', color='#70071c', y=0.995, va='top',
)

plt.savefig(str(FIGURE_PATH), dpi=600, bbox_inches='tight', facecolor='white')
plt.close()
print(f"\nSaved: {FIGURE_PATH}")


# ── Excel ──────────────────────────────────────────────────────────────────────
THIN = Side(style='thin', color='CCCCCC')
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'WeekByWeek_DoseResponse'

ws.append(['Supplementary: Week-by-Week Dose-Response (Wks 1\u20134 each vs Wk 5+ ref) \u2014 PSM + TV Cox'])
ws['A1'].font = Font(bold=True, size=12, color='70071c')
ws.append(['Week 0 (days 1\u20137) excluded  |  Reference: Week 5+ (days 36\u201390)  |  Max follow-up 365 days'])
ws['A2'].font = Font(italic=True, size=10, color='555555')
ws.append([])

cols = ['Comparison', 'Outcome', 'N', 'Events', 'Event %', 'HR (95% CI)', 'p-value']
hdr = ws.max_row + 1
for ci, col in enumerate(cols, 1):
    c = ws.cell(row=hdr, column=ci, value=col)
    c.font      = Font(bold=True, color='FFFFFF', size=10)
    c.fill      = PatternFill('solid', fgColor='70071c')
    c.alignment = Alignment(horizontal='center', wrap_text=True)
ws.row_dimensions[hdr].height = 28

prev_comp, alt = None, 0
for _, row in df_res.iterrows():
    comp = row['Comparison']
    if comp != prev_comp:
        ri = ws.max_row + 1
        for ci in range(1, len(cols) + 1):
            c = ws.cell(row=ri, column=ci, value=comp if ci == 1 else '')
            c.font      = Font(bold=True, color='FFFFFF', size=10)
            c.fill      = PatternFill('solid', fgColor='ba0c2f')
            c.alignment = Alignment(horizontal='left' if ci == 1 else 'center',
                                    vertical='center')
        prev_comp = comp
        alt = 0
    alt += 1
    ri = ws.max_row + 1
    ws.row_dimensions[ri].height = 16
    hr_str = (f"{row['HR']:.2f}  [{row['CI_low']:.2f}-{row['CI_high']:.2f}]"
              if pd.notna(row['HR']) else 'N/A')
    for ci, val in enumerate(['', row['Outcome'], f"{row['N']:,}", f"{row['Events']:,}",
                               f"{row['Event_pct']:.1f}%", hr_str, row['p_str']], 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.fill      = PatternFill('solid', fgColor='fdf5f6') if alt % 2 == 0 else PatternFill()
        c.alignment = Alignment(horizontal='left', vertical='center')
        c.border    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        if ci == 6 and pd.notna(row.get('p_raw')) and row['p_raw'] < 0.05:
            c.font = Font(bold=True, size=10, color='ba0c2f')

for col_ltr, w in zip('ABCDEFG', [30, 22, 10, 10, 10, 24, 12]):
    ws.column_dimensions[col_ltr].width = w

wb.save(str(TABLE_PATH))
print(f"Saved: {TABLE_PATH}")
